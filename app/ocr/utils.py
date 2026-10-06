import os
import boto3
import pdf2image
import requests
import shutil
import subprocess
import tempfile
import time
import logging
import logging.handlers
from urllib.parse import urlparse

import re
import psycopg2
from pypdf import PdfReader
from PIL import Image, ImageEnhance, ImageFilter
from typing import Optional
import cv2
import numpy as np
import multiprocessing

from dotenv import load_dotenv
load_dotenv()

rabbitmq_user = os.getenv("RABBITMQ_USER", "your_rabbit_user")
rabbitmq_pass = os.getenv("RABBITMQ_PASS", "your_rabbit_password")
rabbitmq_url = "http://localhost:15672/api/queues/%2F/process_queue"

CO_QUAN_LIST = [
    "TỈNH ỦY THÁI BÌNH",
    "ĐẢNG BỘ TỈNH THÁI BÌNH",
    "TỈNH ỦY HƯNG YÊN",
    "ĐẢNG BỘ TỈNH HƯNG YÊN",
    "ĐẢNG BỘ TỈNH BÀ RỊA - VŨNG TÀU",
    "THÀNH ỦY VŨNG TÀU"
]

logger = logging.getLogger(__name__)


class TruncatingRotatingHandler(logging.handlers.RotatingFileHandler):
    def doRollover(self):
        if self.stream:
            self.stream.close()
            self.stream = None
        open(self.baseFilename, "w").close()
        self.stream = self._open()

import unicodedata
def clean_text(text):
    BANG_XOA_DAU = str.maketrans(
        "ÁÀẢÃẠĂẮẰẲẴẶÂẤẦẨẪẬĐÈÉẺẼẸÊẾỀỂỄỆÍÌỈĨỊÓÒỎÕỌÔỐỒỔỖỘƠỚỜỞỠỢÚÙỦŨỤƯỨỪỬỮỰÝỲỶỸỴáàảãạăắằẳẵặâấầẩẫậđèéẻẽẹêếềểễệíìỉĩịóòỏõọôốồổỗộơớờởỡợúùủũụưứừửữựýỳỷỹỵ",
        "A"*17 + "D" + "E"*11 + "I"*5 + "O"*17 + "U"*11 + "Y"*5 + "a"*17 + "d" + "e"*11 + "i"*5 + "o"*17 + "u"*11 + "y"*5
    )
    if not unicodedata.is_normalized("NFC", text):
        text = unicodedata.normalize("NFC", text)
    return text.translate(BANG_XOA_DAU)


def get_template_name(CustomerID, ConfigTemplateID, logger):
    try:
        conn = psycopg2.connect(
            host=os.getenv("POSTGRES_HOST"),
            database=os.getenv("POSTGRES_DB"),
            user=os.getenv("POSTGRES_USER"),
            password=os.getenv("POSTGRES_PW")
        )
        cursor = conn.cursor()
        query = """
        SELECT s.name FROM sampleconfigdocuments s
        WHERE s.id = %s AND s.companyid = %s AND s.status != 99
        LIMIT 1;
        """
        cursor.execute(query, (ConfigTemplateID, CustomerID))
        row = cursor.fetchone()

        template_name = row[0] if row else ""
        logger.info(f"DB query for ({CustomerID}, {ConfigTemplateID}) found template name: '{template_name}'")

        cursor.close()
        conn.close()
        return template_name

    except Exception as e:
        logger.error(f"Error loading template name from database: {e}")
        return ""



def parse_s3_path(s3_url: str) -> tuple[str, str]:
    parsed_url = urlparse(s3_url)
    if parsed_url.scheme != 's3':
        raise ValueError(f"Not an S3 URL: {s3_url}")
    return parsed_url.netloc, parsed_url.path.lstrip('/')


def _s3_client(access_key: str, secret_key: str, endpoint: str):
    return boto3.client(
        's3',
        endpoint_url=endpoint,
        aws_access_key_id=access_key,
        aws_secret_access_key=secret_key,
        region_name=os.getenv("S3_REGION", "us-east-1"),
        use_ssl=os.getenv("S3_SSL", "False").lower() in ('true', '1', 't'),
    )


def list_s3_objects(bucket: str, prefix: str, access_key: str, secret_key: str, endpoint: str) -> dict:
    s3_client = _s3_client(access_key, secret_key, endpoint)
    paginator = s3_client.get_paginator('list_objects_v2')
    kwargs = {'Bucket': bucket}
    if prefix:
        kwargs['Prefix'] = prefix
    result = {}
    for page in paginator.paginate(**kwargs):
        for obj in page.get('Contents', []):
            result[obj['Key']] = obj['ETag'].strip('"')
    return result


def download_from_s3(s3_url: str, local_path: str, access_key: str, secret_key: str, endpoint: str) -> dict:
    try:
        s3_client = _s3_client(access_key, secret_key, endpoint)
        bucket, key = parse_s3_path(s3_url)

        logger.info(f"Downloading from bucket: {bucket}, key: {key}")
        response = s3_client.get_object(Bucket=bucket, Key=key)
        content_type = response.get('ContentType', 'application/octet-stream')
        etag = response.get('ETag', '').strip('"')

        os.makedirs(local_path, exist_ok=True)

        file_name = os.path.basename(key)
        file_path = os.path.join(local_path, file_name)

        with open(file_path, 'wb') as f:
            f.write(response['Body'].read())

        logger.info(f"File downloaded successfully to {file_path}")
        return {"file_path": file_path, "content_type": content_type, "etag": etag, "bucket": bucket, "key": key}
    except Exception as e:
        logger.error(f"Error downloading file from S3: {e}", exc_info=True)
        raise

def convertPDFImage(pdf_file_path: str, output_folder: str) -> list[str]:
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
    try:
        thread_count = multiprocessing.cpu_count()
        print(f"Số logical CPU cores: {thread_count}")
        images = pdf2image.convert_from_path(pdf_file_path,dpi=100, thread_count=thread_count)
        image_paths = []
        for i, image in enumerate(images):
            image_path = os.path.join(output_folder, f"page_{i+1}.jpg")
            image.save(image_path, "JPEG", quality=85, optimize=True)
            image_paths.append(image_path)
        return image_paths
    except Exception as e:
        raise


def get_pdf_metadata(pdf_path: str) -> dict:
    try:
        reader = PdfReader(pdf_path)
        metadata = reader.metadata
        logger.info(f"Extracted metadata from {pdf_path}: {metadata}")
        return metadata
    except Exception as e:
        logger.error(f"Error extracting metadata from PDF {pdf_path}: {e}")
        return {}



import numpy as np
import os
def find_skew(gray, delta=0.2, limit=5):
    angles = np.arange(-limit, limit + delta, delta)
    scores = []

    h, w = gray.shape
    center = (w // 2, h // 2)

    for angle in angles:
        M = cv2.getRotationMatrix2D(center, angle, 1.0)
        rotated = cv2.warpAffine(gray, M, (w, h), flags=cv2.INTER_LINEAR, borderMode=cv2.BORDER_REPLICATE)

        projection = np.sum(rotated, axis=1)

        score = np.var(projection)
        scores.append(score)

    best_angle = angles[np.argmax(scores)]
    return best_angle


def deskew(image):
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    angle = find_skew(gray)

    print("Detected skew angle:", angle)

    (h, w) = image.shape[:2]
    center = (w // 2, h // 2)
    M = cv2.getRotationMatrix2D(center, angle, 1.0)

    rotated = cv2.warpAffine(image, M, (w, h), flags=cv2.INTER_LINEAR, borderMode=cv2.BORDER_REPLICATE)
    return rotated


def send_webhook_sync(url: str, data: dict):
    logger.info(f"Sending webhook to {url}")
    max_retries = 3
    for attempt in range(max_retries):
        try:
            response = requests.post(
                url, json=data, headers={"Content-Type": "application/json"}, timeout=30.0
            )
            response.raise_for_status()
            logger.info(f"Webhook sent successfully to {url}, status: {response.status_code}")
            return True
        except requests.exceptions.RequestException as e:
            logger.error(f"Attempt {attempt + 1} failed sending webhook to {url}: {e}")
            if attempt < max_retries - 1:
                time.sleep(2 ** attempt)
    logger.error(f"Failed to send webhook to {url} after {max_retries} attempts.")
    return False



def variants(name: str):
    return {
        "origin": name,
        "lower": name.lower(),
        "clean": clean_text(name)
    }

def process_text(text: str, list_co_quan: list) -> Optional[str]:
    text_clean = clean_text(text)
    text_lower = text.lower()

    for cq in list_co_quan:
        if text == cq["origin"] or text_lower == cq["lower"] or text_clean == cq["clean"]:
            return text

        if (
            text.startswith(cq["origin"]) or
            text_lower.startswith(cq["lower"]) or
            text_clean.startswith(cq["clean"])
        ):
            text = text[len(cq["origin"]):]
            return text.replace("\n", " ").upper() + " " + cq["origin"]

    text = text.replace("\n", " ")
    if text.startswith(" "):
        return text[1:].upper()
    return text.upper()


def is_box_in_top_left(x1, y1, x2, y2, image_width, image_height, ratio=0.5):
    region_w = image_width * ratio
    region_h = image_height * ratio
    return x2 <= region_w and y2 <= region_h




def is_full_caps_word(word: str) -> bool:
    letters = [c for c in word if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters)


def insert_newline_after_title(text: str) -> str:
    words = text.split()
    current = []
    title = None
    title_end_idx = 0

    for i, w in enumerate(words):
        if is_full_caps_word(w):
            current.append(w)
            title_end_idx = i
        else:
            break

    if not current:
        return text

    title = " ".join(current)

    pos = text.find(title)
    if pos == -1:
        return text

    after_title = pos + len(title)

    if after_title < len(text) and text[after_title] == "\n":
        return text

    if text[after_title:after_title+2] == " \n":
        return text

    return text[:after_title] + "\n" + text[after_title:].lstrip()


import re as _re
from pathlib import Path as _Path
import openpyxl as _openpyxl

_SPELL_DICT_PATH = _Path(__file__).parent / 'spell_dict.xlsx'
_spell_cache: list | None = None
_spell_mtime: float = 0.0


def _make_replacer(correct: str):
    has_upper = any(c.isupper() and c.isalpha() for c in correct)

    def _rep(m: _re.Match) -> str:
        if has_upper:
            return correct
        matched = m.group(0)
        if matched and matched[0].isupper() and correct:
            return correct[0].upper() + correct[1:]
        return correct

    return _rep


def _load_spell_dict() -> list:
    global _spell_cache, _spell_mtime
    try:
        mtime = _SPELL_DICT_PATH.stat().st_mtime
    except FileNotFoundError:
        return []
    if _spell_cache is not None and mtime == _spell_mtime:
        return _spell_cache
    entries = []
    wb = _openpyxl.load_workbook(_SPELL_DICT_PATH, read_only=True, data_only=True)
    ws = wb.active
    first = True
    for row in ws.iter_rows(values_only=True):
        if first:
            first = False
            continue
        if not row or len(row) < 2:
            continue
        wrong = str(row[0]).strip() if row[0] is not None else ''
        correct = str(row[1]).strip() if row[1] is not None else ''
        if not wrong or not correct or wrong.startswith('#'):
            continue
        pat = _re.compile(r'(?<!\w)' + _re.escape(wrong) + r'(?!\w)',
                          _re.IGNORECASE | _re.UNICODE)
        entries.append((pat, _make_replacer(correct)))
    wb.close()
    _spell_cache = entries
    _spell_mtime = mtime
    return entries


def fix_spelling(text: str) -> str:
    if not text:
        return text
    for pat, replacer in _load_spell_dict():
        text = pat.sub(replacer, text)
    return text


_ABBREV_PATTERNS = [
    (_re.compile(r'Hội\s+đồng\s+nhân\s+dân', _re.IGNORECASE | _re.UNICODE), 'HĐND'),
    (_re.compile(r'HỘI\s+ĐỒNG\s+NHÂN\s+DÂN', _re.UNICODE), 'HĐND'),
    (_re.compile(r'Ủy\s+ban\s+nhân\s+dân', _re.IGNORECASE | _re.UNICODE), 'UBND'),
    (_re.compile(r'ỦY\s+BAN\s+NHÂN\s+DÂN', _re.UNICODE), 'UBND'),
    (_re.compile(r'\bubnd\b', _re.IGNORECASE | _re.UNICODE), 'UBND'),
    (_re.compile(r'\bhđnd\b', _re.IGNORECASE | _re.UNICODE), 'HĐND'),
    (_re.compile(r'về\s+việc', _re.IGNORECASE | _re.UNICODE), 'v/v'),
]


def normalize_abbrevs(text: str) -> str:
    for pattern, replacement in _ABBREV_PATTERNS:
        text = pattern.sub(replacement, text)
    return text


def sentence_case(text: str) -> str:
    if not text:
        return text
    s = text.lower()
    return s[0].upper() + s[1:]


_ADMIN_UNIT_RE = _re.compile(
    r'\b(thôn|xóm|xã|phường|quận|huyện|tỉnh|thành phố|thị xã|thị trấn)\s+([^\n,;.(]+)',
    _re.IGNORECASE | _re.UNICODE,
)


def _title_vn(s: str) -> str:
    words = s.split()
    return ' '.join(w[0].upper() + w[1:] if w else w for w in words)


def capitalize_after_admin_units(text: str) -> str:
    def _cap(m: _re.Match) -> str:
        return m.group(1) + ' ' + _title_vn(m.group(2).strip())
    return _ADMIN_UNIT_RE.sub(_cap, text)


_GEO   = {'tỉnh','thành phố','tp','quận','huyện','thị xã','thị trấn','xã','phường','thôn'}
_PHR   = {'sở','phòng','ban','chi cục','cục','tổng cục','vụ','viện','trung tâm',
          'văn phòng','trường','bệnh viện','công ty','tập đoàn','chi nhánh','đội',
          'hội đồng','ủy ban','ban quản lý'}
_ORG   = {'nhà nước','chính phủ','quốc hội','bộ','thứ','cơ quan'}
_LEGAL = {'tnhh','cổ phần','hợp danh','tư nhân','một thành viên'}
_ABBR  = {'ubnd','hđnd','tnhh','cp','tnhh-mtv','bql','ban ql'}
_GUARD = {'và','or','&'}
_CONJ  = {'và','hoặc','or','&','-','–'}

_COMPOUNDS = [
    ('chi cục', 'chi~cục'), ('tổng cục', 'tổng~cục'), ('thành phố', 'thành~phố'),
    ('thị xã', 'thị~xã'), ('thị trấn', 'thị~trấn'), ('ban quản lý', 'ban~quản~lý'),
    ('trung tâm', 'trung~tâm'), ('văn phòng', 'văn~phòng'), ('bệnh viện', 'bệnh~viện'),
    ('hội đồng', 'hội~đồng'), ('ủy ban', 'ủy~ban'), ('chi nhánh', 'chi~nhánh'),
    ('một thành viên', 'một~thành~viên'), ('tập đoàn', 'tập~đoàn'),
]

_ALREADY_UPPER_RE = _re.compile(r'^[A-ZĐÁÀẢÃẠĂẮẰẲẴẶÂẤẦẨẪẬÉÈẺẼẸÊẾỀỂỄỆÍÌỈĨỊÓÒỎÕỌÔỐỒỔỖỘƠỚỜỞỠỢÚÙỦŨỤƯỨỪỬỮỰÝỲỶỸỴ]{2,}', _re.UNICODE)
_STARTS_UPPER_RE  = _re.compile(r'^[A-ZĐÁÀẢÃẠĂẮẰẲẴẶÂẤẦẨẪẬÉÈẺẼẸÊẾỀỂỄỆÍÌỈĨỊÓÒỎÕỌÔỐỒỔỖỘƠỚỜỞỠỢÚÙỦŨỤƯỨỪỬỮỰÝỲỶỸỴ]', _re.UNICODE)


def _vn_hoa_dau(s: str) -> str:
    return s[0].upper() + s[1:] if s else s


_VN_TEN_RIENG = [
    (_re.compile(r'(?<!\w)ha noi(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Hà Nội'),
    (_re.compile(r'(?<!\w)ho chi minh(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Hồ Chí Minh'),
    (_re.compile(r'(?<!\w)da nang(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Đà Nẵng'),
    (_re.compile(r'(?<!\w)hai phong(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Hải Phòng'),
    (_re.compile(r'(?<!\w)can tho(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Cần Thơ'),
    (_re.compile(r'(?<!\w)đak lak(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Đắk Lắk'),
    (_re.compile(r'(?<!\w)đăk lăk(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Đắk Lắk'),
    (_re.compile(r'(?<!\w)lang son(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Lạng Sơn'),
    (_re.compile(r'(?<!\w)ninh thuan(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Ninh Thuận'),
    (_re.compile(r'(?<!\w)binh duong(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Bình Dương'),
    (_re.compile(r'(?<!\w)khanh hoa(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Khánh Hòa'),
    (_re.compile(r'(?<!\w)quang nam(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Quảng Nam'),
    (_re.compile(r'(?<!\w)thua thien hue(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Thừa Thiên Huế'),
    (_re.compile(r'(?<!\w)lam dong(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Lâm Đồng'),
    (_re.compile(r'(?<!\w)dong nai(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Đồng Nai'),
    (_re.compile(r'(?<!\w)dong thap(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Đồng Tháp'),
    (_re.compile(r'(?<!\w)tien giang(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Tiền Giang'),
    (_re.compile(r'(?<!\w)vung tau(?!\w)', _re.IGNORECASE | _re.UNICODE), 'Vũng Tàu'),
]


def _vn_ten_rieng(s: str) -> str:
    for pat, repl in _VN_TEN_RIENG:
        s = pat.sub(repl, s)
    return s


def chuan_hoa_tac_gia(text: str) -> str:
    if not text:
        return text
    s = _re.sub(r'\s+', ' ', text.replace('\n', ' ')).strip()
    if not s:
        return s
    s = s.lower()
    s = _re.sub(r'(ủy|uỷ)\s+ban\s+nhân\s+dân', 'UBND', s)
    s = _re.sub(r'hội\s+đồng\s+nhân\s+dân',     'HĐND', s)
    for src, dst in _COMPOUNDS:
        s = s.replace(src, dst)
    parts = [w for w in s.split(' ') if w]
    n, outp, mode, cap, tcnt, prev_k = len(parts), [], 'low', False, 0, ''
    for w in parts:
        k = w.strip('.,;:()[]"\'-')
        if _ALREADY_UPPER_RE.match(w):
            outp.append(w); mode = 'low'
        elif k in _ABBR:
            outp.append(w.upper())
            if mode != 'phrase': mode = 'low'; cap = False
        elif k == 'tp':
            outp.append('TP'); mode = 'title'; cap = False; tcnt = 0
        elif k in _LEGAL:
            outp.append(w); mode = 'phrase'; cap = True
        elif k in _GEO:
            outp.append(w); mode = 'title'; cap = False; tcnt = 0
        elif k in _ORG:
            outp.append(_vn_hoa_dau(w) if (mode == 'phrase' and cap) else w)
            mode = 'low'; cap = False
        elif k in _PHR:
            outp.append(w if prev_k in _GUARD else _vn_hoa_dau(w))
            mode = 'phrase'; cap = True
        elif k in _CONJ or k == ',':
            outp.append(w)
            if mode == 'phrase': cap = True
        elif mode == 'title' and tcnt < 4:
            outp.append(_vn_hoa_dau(w)); tcnt += 1
        elif mode == 'phrase' and cap:
            outp.append(_vn_hoa_dau(w)); cap = False
        else:
            if mode == 'title': mode = 'low'
            outp.append(w)
        if mode == 'phrase' and n <= 12 and w and w[-1] in ',-–/':
            cap = True
        prev_k = k
    s = ' '.join(outp).replace('~', ' ')
    if s and not _STARTS_UPPER_RE.match(s):
        s = _vn_hoa_dau(s)
    return _vn_ten_rieng(s)


_SENSITIVE_WORDS_PATH = _Path(__file__).parent / 'sensitive_words.txt'
_sensitive_cache: list | None = None
_sensitive_mtime: float = 0.0


def _load_sensitive_words() -> list:
    global _sensitive_cache, _sensitive_mtime
    try:
        mtime = _SENSITIVE_WORDS_PATH.stat().st_mtime
    except FileNotFoundError:
        return []
    if _sensitive_cache is not None and mtime == _sensitive_mtime:
        return _sensitive_cache
    words = []
    with open(_SENSITIVE_WORDS_PATH, 'r', encoding='utf-8') as f:
        for line in f:
            word = line.strip()
            if not word or word.startswith('#'):
                continue
            words.append(word)
    words.sort(key=len, reverse=True)
    patterns = []
    for word in words:
        body = r'\s+'.join(_re.escape(part) for part in word.split())
        patterns.append(_re.compile(r'(?<!\w)' + body + r'(?!\w)', _re.IGNORECASE | _re.UNICODE))
    _sensitive_cache = patterns
    _sensitive_mtime = mtime
    return patterns


def mask_sensitive_words(text: str) -> str:
    if not text:
        return text
    for pat in _load_sensitive_words():
        text = pat.sub(lambda m: '*' * len(m.group(0)), text)
    return text


TEXT_DOC_EXTS = {".txt", ".docx", ".doc"}


def extract_document_text(file_path: str, temp_dir: str = None) -> str:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def _row_cells_text(row) -> list:
        texts, prev_tc = [], None
        for cell in row.cells:
            if cell._tc is prev_tc:
                continue
            prev_tc = cell._tc
            texts.append(cell.text.strip())
        return texts

    def _read_docx(path: str) -> str:
        document = docx.Document(path)
        parts = []
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                text = Paragraph(child, document).text
                if text.strip():
                    parts.append(text)
            elif child.tag == qn("w:tbl"):
                for row in Table(child, document).rows:
                    cells = _row_cells_text(row)
                    if any(cells):
                        parts.append(" | ".join(cells))
        return "\n".join(parts)

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        for enc in ("utf-8", "utf-8-sig", "cp1258", "latin-1"):
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, LookupError):
                continue
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    if ext == ".docx":
        return _read_docx(file_path)

    if ext == ".doc":
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise ValueError(
                "Không xử lý được file .doc: cần cài LibreOffice (lệnh `soffice` trong PATH) "
                "trên server để convert .doc -> .docx trước khi trích xuất text."
            )
        work_dir = temp_dir or tempfile.mkdtemp(prefix="doc_convert_")
        os.makedirs(work_dir, exist_ok=True)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", work_dir, file_path],
            check=True, capture_output=True, timeout=120,
        )
        stem = os.path.splitext(os.path.basename(file_path))[0]
        out_path = os.path.join(work_dir, f"{stem}.docx")
        if not os.path.exists(out_path):
            raise ValueError(f"Convert .doc -> .docx thất bại: không tìm thấy {out_path}")
        return _read_docx(out_path)

    raise ValueError(f"Định dạng văn bản không hỗ trợ: {ext}")
